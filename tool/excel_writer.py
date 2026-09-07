"""生成 Excel：明细表 + 汇总表。"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from tool.parser import LEVELS
from tool.taskfiles import ServerResult

DETAIL_HEADERS = ["IP", "主机名", "告警序号", "级别", "日期", "时间", "告警内容"]
SUMMARY_HEADERS = ["IP", "主机名", "状态", "失败原因", "告警总数"] + list(LEVELS)  # 末尾追加各告警级别计数列
STATUS_CN = {"SUCCESS": "成功", "FAIL": "失败", "STOPPED": "已停止"}  # 汇总表状态列的中文映射


def write_workbook(path: Path, results: list[ServerResult], sim_mode: bool = False) -> None:
    """把采集结果写入 Excel 工作簿。

    结构：明细表（每台设备的每条告警一行）+ 汇总表（每台设备一行，含状态、
    失败原因、告警总数与各级别计数）。sim_mode 时追加"说明"表标注模拟数据，
    防止模拟产物被误当真实结果外发。目录不存在时自动创建。
    """
    wb = Workbook()

    ws_detail = wb.active
    ws_detail.title = "明细"
    ws_detail.append(DETAIL_HEADERS)
    for r in results:
        for a in r.alarms:
            ws_detail.append(
                [r.ip, r.hostname, a["index"], a["level"], a["date"], a["time"], a["info"]]
            )
    _style_header(ws_detail)
    _auto_width(ws_detail)

    ws_summary = wb.create_sheet("汇总")
    ws_summary.append(SUMMARY_HEADERS)
    for r in results:
        # Counter 按出现次数统计各级别；按 LEVELS 固定顺序取值，
        # 保证列序稳定且未出现的级别补 0
        counts = Counter(a["level"] for a in r.alarms)
        ws_summary.append(
            [r.ip, r.hostname, STATUS_CN.get(r.status, r.status), r.reason, len(r.alarms)]
            + [counts.get(lv, 0) for lv in LEVELS]
        )
    _style_header(ws_summary)
    _auto_width(ws_summary)

    if sim_mode:
        ws_note = wb.create_sheet("说明")
        ws_note.append(["注意：本文件为模拟数据（测试用），非真实设备结果。"])
        ws_note.column_dimensions["A"].width = 60

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _style_header(ws) -> None:
    """表头行加粗 + 浅蓝底色，与数据行区分。"""
    fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill


def _auto_width(ws) -> None:
    """按本列最长内容自动设列宽，上限 60 防止超长告警把列撑爆。"""
    for col in ws.columns:
        width = max(_display_width(c.value) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 60)


def _display_width(value) -> int:
    """计算文本显示宽度：CJK 字符在 Excel 中约占 2 个宽度单位，其余按 1。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(value))
