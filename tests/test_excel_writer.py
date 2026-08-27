from pathlib import Path

from openpyxl import load_workbook

from tool.excel_writer import write_workbook
from tool.taskfiles import ServerResult


def make_result(ip: str, status: str, reason: str, alarms: list[dict]) -> ServerResult:
    return ServerResult(ip=ip, hostname=f"host-{ip}", status=status,
                        reason=reason, raw_path=None, alarms=alarms)


def test_workbook_has_two_sheets_with_correct_rows(tmp_path: Path):
    results = [
        make_result("1.1.1.1", "SUCCESS", "", [
            {"index": "1", "level": "Critical", "date": "2023-04-28",
             "time": "01:02:41+08:00", "info": "POWER failed"},
            {"index": "2", "level": "Minor", "date": "2023-04-29",
             "time": "09:00:00+08:00", "info": "Fan speed abnormal"},
        ]),
        make_result("2.2.2.2", "FAIL", "连接失败", []),
    ]
    out = tmp_path / "结果.xlsx"
    write_workbook(out, results)

    wb = load_workbook(out)
    assert wb.sheetnames == ["明细", "汇总"]

    ws = wb["明细"]
    rows = list(ws.values)
    assert rows[0] == ("IP", "主机名", "告警序号", "级别", "日期", "时间", "告警内容")
    assert len(rows) == 3  # 表头 + 2 条告警
    assert rows[1][0] == "1.1.1.1" and rows[1][3] == "Critical"
    assert rows[2][6] == "Fan speed abnormal"

    ws2 = wb["汇总"]
    rows2 = list(ws2.values)
    assert rows2[0] == ("IP", "主机名", "状态", "失败原因", "告警总数",
                        "Critical", "Major", "Minor", "Warning")
    assert len(rows2) == 3  # 表头 + 2 台
    # 注：openpyxl 保存时丢弃空字符串，空单元格读回为 None
    assert rows2[1][:5] == ("1.1.1.1", "host-1.1.1.1", "成功", None, 2)
    assert rows2[1][5] == 1 and rows2[1][7] == 1  # Critical 1、Minor 1（列 5-8 为四级计数）
    assert rows2[2][:4] == ("2.2.2.2", "host-2.2.2.2", "失败", "连接失败")


def test_stopped_status_shown_in_summary(tmp_path: Path):
    results = [make_result("3.3.3.3", "STOPPED", "用户停止", [])]
    out = tmp_path / "停止.xlsx"
    write_workbook(out, results)
    wb = load_workbook(out)
    ws = wb["汇总"]
    rows = list(ws.values)
    assert rows[1][2] == "已停止"


def test_all_level_counts_in_summary(tmp_path: Path):
    results = [make_result("4.4.4.4", "SUCCESS", "", [
        {"index": "1", "level": "Critical", "date": "d", "time": "t", "info": "i"},
        {"index": "2", "level": "Major", "date": "d", "time": "t", "info": "i"},
        {"index": "3", "level": "Minor", "date": "d", "time": "t", "info": "i"},
        {"index": "4", "level": "Warning", "date": "d", "time": "t", "info": "i"},
    ])]
    out = tmp_path / "levels.xlsx"
    write_workbook(out, results)
    rows = list(load_workbook(out)["汇总"].values)
    assert rows[1][5:9] == (1, 1, 1, 1)


def test_sim_mode_adds_note_sheet(tmp_path: Path):
    out = tmp_path / "sim.xlsx"
    write_workbook(out, [make_result("1.1.1.1", "SUCCESS", "", [])], sim_mode=True)
    wb = load_workbook(out)
    assert wb.sheetnames == ["明细", "汇总", "说明"]
    assert "模拟数据" in wb["说明"].cell(1, 1).value


def test_real_mode_has_no_note_sheet(tmp_path: Path):
    out = tmp_path / "real.xlsx"
    write_workbook(out, [make_result("1.1.1.1", "SUCCESS", "", [])], sim_mode=False)
    wb = load_workbook(out)
    assert wb.sheetnames == ["明细", "汇总"]
